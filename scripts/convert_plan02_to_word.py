from pathlib import Path
import re
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'docs' / 'PLAN-02-backend-sections.docx'
SOURCES = [
    ROOT / 'docs' / 'PLAN-02-backend-sections.md',
    ROOT / 'docs' / 'PLAN-02-backend-evidence.md',
]


def add_inline_text(paragraph, text):
    parts = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
        else:
            paragraph.add_run(part)


def add_markdown(document, markdown):
    in_code = False
    code_lines = []
    list_style = None
    for raw_line in markdown.splitlines():
        line = raw_line.rstrip()
        if line.startswith('```'):
            if in_code:
                paragraph = document.add_paragraph(style='Code Block')
                paragraph.add_run('\n'.join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not line.strip():
            list_style = None
            continue
        heading = re.match(r'^(#{1,4})\s+(.*)$', line)
        if heading:
            level = min(len(heading.group(1)), 3)
            document.add_heading(heading.group(2).strip(), level=level)
            list_style = None
            continue
        if line.startswith('|'):
            if set(line.replace('|', '').replace('-', '').replace(':', '').strip()) == set():
                continue
            cells = [cell.strip() for cell in line.strip('|').split('|')]
            table = document.add_table(rows=1, cols=len(cells))
            table.style = 'Table Grid'
            for cell, value in zip(table.rows[0].cells, cells):
                cell.text = value
            list_style = None
            continue
        bullet = re.match(r'^\s*[-*]\s+(.*)$', line)
        numbered = re.match(r'^\s*\d+\.\s+(.*)$', line)
        if bullet or numbered:
            style = 'List Bullet' if bullet else 'List Number'
            paragraph = document.add_paragraph(style=style)
            add_inline_text(paragraph, (bullet or numbered).group(1))
            list_style = style
            continue
        paragraph = document.add_paragraph()
        add_inline_text(paragraph, line)
        list_style = None


def main():
    document = Document()
    styles = document.styles
    if 'Code Block' not in [style.name for style in styles]:
        code_style = styles.add_style('Code Block', WD_STYLE_TYPE.PARAGRAPH)
        code_style.font.name = 'Consolas'
        code_style.font.size = Pt(8)
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    document.add_heading('PLAN-02 Backend Sections', level=0)
    document.add_paragraph('PickSell Marketplace and Logistics | Prepared for Justin Paul M. Plantilla')
    for index, source in enumerate(SOURCES):
        if index:
            document.add_page_break()
        add_markdown(document, source.read_text(encoding='utf-8'))
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    main()
